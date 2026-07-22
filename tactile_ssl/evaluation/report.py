"""CSV and XLSX reports for downstream significance results."""

from __future__ import annotations

import csv
import math
from pathlib import Path
from typing import Any, Dict, List, Mapping, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TASK_LAYOUTS: Mapping[str, Tuple[str, Sequence[Tuple[str, Sequence[str]]]]] = {
    "force": (
        "Force Estimation",
        (
            ("rmse", ("base (signal+pos)", "only signal")),
            ("rmse_x", ("base (signal+pos)", "only signal", "paper result (approx.)")),
            ("rmse_y", ("base (signal+pos)", "only signal", "paper result (approx.)")),
            ("rmse_z", ("base (signal+pos)", "only signal", "paper result (approx.)")),
        ),
    ),
    "pose": (
        "Pose Estimation",
        tuple(
            (metric, ("base (signal+pos)", "only signal", "paper result (approx.)"))
            for metric in ("rmse_x", "rmse_y", "rmse_theta", "acc_x", "acc_y", "acc_theta")
        ),
    ),
    "object_classification": (
        "Object Classification",
        (("acc", ("base (signal+pos)", "only signal", "paper result (approx.)")),),
    ),
}


CSV_FIELDS = (
    "task",
    "comparison_block_id",
    "comparison_block_name",
    "baseline_id",
    "baseline_name",
    "statistics_cache_key",
    "statistics_cache_hit",
    "experiment_id",
    "method",
    "is_baseline",
    "variant",
    "metric",
    "estimate",
    "bootstrap_se",
    "delta",
    "delta_se",
    "improvement",
    "raw_pvalue",
    "checkpoint",
    "artifact_path",
    "num_examples",
    "num_groups",
    "dataset_fingerprint",
    "seed",
)


def _number(value: Any) -> str:
    return f"{float(value):.4f}"


def _estimate_text(row: Mapping[str, Any]) -> str:
    return f"{_number(row['estimate'])} ± {_number(row['bootstrap_se'])}"


def _delta_text(row: Mapping[str, Any]) -> str:
    improvement = row.get("improvement")
    improvement_text = (
        "n/a"
        if improvement is None or not math.isfinite(float(improvement))
        else f"{100.0 * float(improvement):+.1f}%"
    )
    return f"{_number(row['delta'])} ± {_number(row['delta_se'])} ({improvement_text})"


def _pvalue_text(value: float) -> str:
    return "<0.0001" if value < 0.0001 else f"{value:.4f}"


def write_csv(path: Path, rows: Sequence[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _style_presentation_sheet(sheet, max_column: int) -> None:
    navy = "1F4E78"
    blue = "D9EAF7"
    white = "FFFFFF"
    thin = Side(style="thin", color="B7C9D6")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in sheet[2]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=max_column):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if cell.row > 2:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "B3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(max_column)}{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:2"
    sheet.column_dimensions["A"].width = 34
    for column in range(2, max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 22
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 44


def _write_task_sheet(sheet, task: str, rows: Sequence[Mapping[str, Any]]) -> None:
    title, layout = TASK_LAYOUTS[task]
    sheet.cell(1, 1, title)
    sheet.cell(2, 1, "method name")
    metric_columns: Dict[Tuple[str, str], int] = {}
    column = 2
    for metric, variants in layout:
        start = column
        for variant in variants:
            sheet.cell(2, column, variant)
            metric_columns[(metric, variant)] = column
            column += 1
        end = column - 1
        sheet.cell(1, start, metric)
        if end > start:
            sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)

    by_block: Dict[str, List[Mapping[str, Any]]] = {}
    block_order: List[str] = []
    for row in rows:
        block_id = str(row.get("comparison_block_id", "default"))
        if block_id not in by_block:
            block_order.append(block_id)
            by_block[block_id] = []
        by_block[block_id].append(row)

    output_row = 3
    block_header_rows: List[int] = []
    baseline_rows: List[int] = []
    method_rows: List[int] = []
    delta_rows: List[int] = []
    pvalue_rows: List[int] = []
    for block_id in block_order:
        block_rows = by_block[block_id]
        block_name = str(block_rows[0].get("comparison_block_name", ""))
        if block_name:
            sheet.cell(output_row, 1, block_name)
            sheet.merge_cells(
                start_row=output_row,
                start_column=1,
                end_row=output_row,
                end_column=column - 1,
            )
            block_header_rows.append(output_row)
            output_row += 1

        by_experiment: Dict[str, List[Mapping[str, Any]]] = {}
        experiment_order: List[str] = []
        for row in block_rows:
            experiment_id = str(row["experiment_id"])
            if experiment_id not in by_experiment:
                experiment_order.append(experiment_id)
                by_experiment[experiment_id] = []
            by_experiment[experiment_id].append(row)

        baseline_name = str(
            block_rows[0].get("baseline_name")
            or next(row["method"] for row in block_rows if row["is_baseline"])
        )
        for experiment_id in experiment_order:
            experiment_rows = by_experiment[experiment_id]
            is_baseline = bool(experiment_rows[0]["is_baseline"])
            labels = [str(experiment_rows[0]["method"])]
            if not is_baseline:
                labels.extend((f"Δ vs {baseline_name} (improvement)", "p-value"))
            for offset, label in enumerate(labels):
                sheet.cell(output_row + offset, 1, label)
            for row in experiment_rows:
                target_column = metric_columns[(str(row["metric"]), str(row["variant"]))]
                sheet.cell(output_row, target_column, _estimate_text(row))
                if not is_baseline:
                    sheet.cell(output_row + 1, target_column, _delta_text(row))
                    sheet.cell(output_row + 2, target_column, _pvalue_text(float(row["raw_pvalue"])))
            if is_baseline:
                baseline_rows.append(output_row)
            else:
                method_rows.append(output_row)
                delta_rows.append(output_row + 1)
                pvalue_rows.append(output_row + 2)
            output_row += len(labels)
    _style_presentation_sheet(sheet, column - 1)
    row_styles = (
        (block_header_rows, "FFDED4ED", True),
        (baseline_rows, "FFD9EAF7", True),
        (method_rows, "FFE2F0D9", False),
        (delta_rows, "FFFCE4D6", False),
        (pvalue_rows, "FFFFF2CC", False),
    )
    for row_indices, color, bold in row_styles:
        for row_index in row_indices:
            for cell in sheet[row_index][: column - 1]:
                cell.fill = PatternFill("solid", fgColor=color)
                if bold:
                    cell.font = Font(bold=True)
            sheet.cell(row_index, 1).alignment = Alignment(vertical="center", wrap_text=True)


def _write_means_task_sheet(sheet, task: str, rows: Sequence[Mapping[str, Any]]) -> None:
    """Write the legacy presentation layout with one mean-only row per method."""
    title, layout = TASK_LAYOUTS[task]
    sheet.cell(1, 1, title)
    sheet.cell(2, 1, "method name")
    metric_columns: Dict[Tuple[str, str], int] = {}
    column = 2
    for metric, variants in layout:
        start = column
        for variant in variants:
            sheet.cell(2, column, variant)
            metric_columns[(metric, variant)] = column
            column += 1
        end = column - 1
        sheet.cell(1, start, metric)
        if end > start:
            sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)

    by_block: Dict[str, List[Mapping[str, Any]]] = {}
    block_order: List[str] = []
    for row in rows:
        block_id = str(row.get("comparison_block_id", "default"))
        if block_id not in by_block:
            block_order.append(block_id)
            by_block[block_id] = []
        by_block[block_id].append(row)

    output_row = 3
    block_header_rows: List[int] = []
    baseline_rows: List[int] = []
    method_rows: List[int] = []
    for block_id in block_order:
        block_rows = by_block[block_id]
        block_name = str(block_rows[0].get("comparison_block_name", ""))
        if block_name:
            sheet.cell(output_row, 1, block_name)
            sheet.merge_cells(
                start_row=output_row,
                start_column=1,
                end_row=output_row,
                end_column=column - 1,
            )
            block_header_rows.append(output_row)
            output_row += 1

        by_experiment: Dict[str, List[Mapping[str, Any]]] = {}
        experiment_order: List[str] = []
        for row in block_rows:
            experiment_id = str(row["experiment_id"])
            if experiment_id not in by_experiment:
                experiment_order.append(experiment_id)
                by_experiment[experiment_id] = []
            by_experiment[experiment_id].append(row)

        for experiment_id in experiment_order:
            experiment_rows = by_experiment[experiment_id]
            is_baseline = bool(experiment_rows[0]["is_baseline"])
            sheet.cell(output_row, 1, str(experiment_rows[0]["method"]))
            for row in experiment_rows:
                target_column = metric_columns[(str(row["metric"]), str(row["variant"]))]
                cell = sheet.cell(output_row, target_column, float(row["estimate"]))
                cell.number_format = "0.0000"
            (baseline_rows if is_baseline else method_rows).append(output_row)
            output_row += 1

    _style_presentation_sheet(sheet, column - 1)
    for row_indices, color, bold in (
        (block_header_rows, "FFDED4ED", True),
        (baseline_rows, "FFD9EAF7", True),
        (method_rows, "FFE2F0D9", False),
    ):
        for row_index in row_indices:
            for cell in sheet[row_index][: column - 1]:
                cell.fill = PatternFill("solid", fgColor=color)
                if bold:
                    cell.font = Font(bold=True)
            sheet.cell(row_index, 1).alignment = Alignment(vertical="center", wrap_text=True)


def _write_flat_sheet(sheet, rows: Sequence[Mapping[str, Any]], fields: Sequence[str]) -> None:
    sheet.append(list(fields))
    for row in rows:
        sheet.append([row.get(field) for field in fields])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    for index, field in enumerate(fields, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(45, max(14, len(field) + 2))


def write_xlsx(
    path: Path,
    rows: Sequence[Mapping[str, Any]],
    experiments: Sequence[Mapping[str, Any]],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for task in TASK_LAYOUTS:
        task_rows = [row for row in rows if row["task"] == task]
        if task_rows:
            title = TASK_LAYOUTS[task][0]
            _write_task_sheet(workbook.create_sheet(title), task, task_rows)
    _write_flat_sheet(workbook.create_sheet("Statistics"), rows, CSV_FIELDS)
    experiment_fields = (
        "task",
        "comparison_block_id",
        "comparison_block_name",
        "baseline_id",
        "baseline_name",
        "experiment_id",
        "method",
        "is_baseline",
        "variant",
        "directory",
        "checkpoint",
        "artifact_path",
        "config_path",
        "dataset_fingerprint",
        "num_examples",
        "num_groups",
        "seed",
    )
    _write_flat_sheet(workbook.create_sheet("Experiments"), experiments, experiment_fields)
    workbook.save(path)


def write_means_xlsx(path: Path, rows: Sequence[Mapping[str, Any]]) -> None:
    """Write mean estimates in the legacy task-table layout."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for task in TASK_LAYOUTS:
        task_rows = [row for row in rows if row["task"] == task]
        if task_rows:
            title = TASK_LAYOUTS[task][0]
            _write_means_task_sheet(workbook.create_sheet(title), task, task_rows)
    workbook.save(path)


def write_reports(
    output_dir: Path,
    rows: Sequence[Mapping[str, Any]],
    experiments: Sequence[Mapping[str, Any]],
) -> Tuple[Path, Path]:
    output_dir = Path(output_dir)
    csv_path = output_dir / "significance_results.csv"
    xlsx_path = output_dir / "significance_report.xlsx"
    means_xlsx_path = output_dir / "downstream_means.xlsx"
    write_csv(csv_path, rows)
    write_xlsx(xlsx_path, rows, experiments)
    write_means_xlsx(means_xlsx_path, rows)
    return xlsx_path, csv_path
